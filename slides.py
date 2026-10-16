"""Один шаблон PPTX с редактируемыми таблицами, диаграммами и книгами."""

from datetime import datetime, timezone
from decimal import Decimal, ROUND_CEILING, ROUND_FLOOR, localcontext
from io import BytesIO
from typing import Mapping, Optional, Sequence, Tuple

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.chart.chart import Chart
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION,
                             XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK)
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE
from pptx.opc.oxml import serialize_part_xml
from pptx.oxml import parse_xml
from pptx.oxml.xmlchemy import OxmlElement
from pptx.presentation import Presentation as PresentationDocument
from pptx.slide import Slide
from pptx.util import Inches, Pt
from xlsxwriter import Workbook

import data as report_data
from data import Metric, RATIO_NOTE, TABLE_HEADER, group_metrics, table_rows, validate_report


TEXT = '111827'
PLAN = 'CBD5E1'
ACTUAL = '2563EB'
WIDTH = 12192000
CONTENT_WIDTH = 12 + 1 / 3
FONT_NAME = 'Arial'
PRODUCT_NAME = 'Report Slides'
PROPERTIES_NAMESPACE = '{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}'
AXIS_PADDING = Decimal('0.3')
AXIS_TARGET_INTERVALS = 4
AXIS_MIN_INTERVALS = 3
AXIS_MIN_STEP = Decimal('0.01')
AXIS_STEPS = (Decimal('1'), Decimal('2'), Decimal('5'), Decimal('10'))
CHART_HEIGHT = Decimal('5.1')
CHART_TOP = Decimal('1.15')
CHART_GUTTER = Decimal('4.05')
CONTENT_MARGINS = Decimal(1)
AXIS_FONT_SIZE = Decimal(14)
POINTS_PER_INCH = Decimal(72)
AXIS_DIGIT_WIDTH = Decimal('0.56')
AXIS_SIGN_WIDTH = Decimal('0.34')
AXIS_SEPARATOR_WIDTH = Decimal('0.28')
AXIS_MIN_MARGIN = Decimal('0.18')
AXIS_TICK_MARGIN = Decimal('0.06')
AXIS_VALUE_MARGIN = Decimal('0.12')
AXIS_TOP_MARGIN = Decimal('0.03')
AXIS_BOTTOM_MARGIN = Decimal('0.17')
AXIS_LEGEND_MARGIN = Decimal('0.45')
AXIS_LABEL_ROTATION = -2700000
AXIS_NOTE_TOP = Decimal('6.34')
AXIS_NOTE_HEIGHT = Decimal('0.27')
CATEGORY_LINE_HEIGHT = 15


class _SafeChartData(CategoryChartData):
    @property
    def xlsx_blob(self):
        # Штатный writer превращает =текст и URL в формулы и ссылки.
        output = BytesIO()
        with Workbook(output, {'in_memory': True, 'strings_to_formulas': False,
                               'strings_to_urls': False}) as book:
            sheet = book.add_worksheet('Sheet1')
            for row, category in enumerate(self.categories, 1):
                sheet.write_string(row, 0, category.label)
            for column, series in enumerate(self, 1):
                sheet.write_string(0, column, series.name)
                number_format = book.add_format({'num_format': series.number_format})
                for row, value in enumerate(series.values, 1):
                    sheet.write_number(row, column, value, number_format)
        return output.getvalue()


def slide_pages(metrics):
    tables = tuple(metrics[start:start + 4] for start in range(0, len(metrics), 4))
    charts = tuple((unit, group[start:start + 4]) for unit, group in group_metrics(metrics)
                   for start in range(0, len(group), 4))
    return tables, charts


def _font(font, size, bold=False, color=TEXT):
    font.name = 'Arial'
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor.from_string(color)


def _frame(frame, text, size, bold=False):
    frame.text = text.replace('\r\n', '\n').replace('\r', '\n')
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for paragraph in frame.paragraphs:
        _font(paragraph.font, size, bold)
        # Штатная табуляция в полдюйма вытесняет допустимый длинный текст за слайд.
        paragraph._p.get_or_add_pPr().set('defTabSz', str(Pt(size)))
        paragraph.line_spacing = 1.0
        paragraph.space_before = paragraph.space_after = Pt(0)


def _text(slide, text, left, top, width, height, size=16, bold=False):
    frame = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    _frame(frame, text, size, bold)
    return frame


def _new_slide(prs, title, period, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    _text(slide, title, .5, .4, CONTENT_WIDTH, .55, 24, True)
    _text(slide, period, .5, 6.64, 10.8, .46)
    _text(slide, str(number), 11.6, 6.64, 1.2, .46).paragraphs[0].alignment = PP_ALIGN.RIGHT
    return slide


def _table_slide(prs, page, period, number, index, count):
    slide = _new_slide(prs, f'План и факт · таблица {index} из {count}', period, number)
    table = slide.shapes.add_table(len(page) + 1, 6, Inches(.5), Inches(1.05),
                                   Inches(CONTENT_WIDTH), Inches(.75 + 1.13 * len(page))).table
    table.first_row = False
    widths = (3.55, 1.25, 1.7, 1.7, 1.8, CONTENT_WIDTH - 10)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for index, values in enumerate([TABLE_HEADER, *table_rows(page)]):
        table.rows[index].height = Inches(.75 if index == 0 else 1.13)
        for column, value in enumerate(values):
            cell = table.cell(index, column)
            cell.margin_left = cell.margin_right = Inches(.08)
            cell.margin_top = cell.margin_bottom = Inches(.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(PLAN if index == 0 else 'FFFFFF')
            _frame(cell.text_frame, value, 16, index == 0)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT if column < 2 else PP_ALIGN.RIGHT
    _text(slide, RATIO_NOTE, .5, 6.34, CONTENT_WIDTH, .27)


def _plot_layout(chart: Chart, coordinates: Optional[Tuple[Decimal, Decimal, Decimal, Decimal]] = None) -> None:
    # Явная область построения выравнивает график с отдельными текстовыми подписями.
    plot_area = chart._chartSpace.chart.plotArea
    layout = OxmlElement('c:layout')
    manual = OxmlElement('c:manualLayout')
    positions = (('x', '.08'), ('y', '.03'), ('w', '.86'), ('h', '.80'))
    position_mode = 'factor'
    if coordinates is not None:
        positions = tuple((tag, str(value)) for tag, value in zip(('x', 'y', 'w', 'h'), coordinates))
        position_mode = 'edge'
    for tag, value in (('layoutTarget', 'inner'), ('xMode', position_mode), ('yMode', position_mode),
                       ('wMode', 'factor'), ('hMode', 'factor')) + positions:
        element = OxmlElement('c:' + tag)
        element.set('val', value)
        manual.append(element)
    layout.append(manual)
    old = plot_area.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}layout')
    if old is not None:
        plot_area.remove(old)
    plot_area.insert(0, layout)


def _calculate_axis_scale(values: Sequence[Decimal]) -> Tuple[Decimal, Decimal, Decimal]:
    with localcontext() as context:
        context.prec = 28
        low = min(Decimal(0), *values)
        high = max(Decimal(0), *values)
        if low == high:
            return Decimal(0), Decimal(1), Decimal(1)
        padding = (high - low) * AXIS_PADDING
        minimum = low - padding if low < 0 else Decimal(0)
        maximum = high + padding if high > 0 else Decimal(0)
        target = (maximum - minimum) / AXIS_TARGET_INTERVALS
        magnitude = Decimal(10) ** target.adjusted()
        step = next(multiplier * magnitude for multiplier in AXIS_STEPS if multiplier * magnitude >= target)
        step = max(AXIS_MIN_STEP, step)
        minimum = (minimum / step).to_integral_value(rounding=ROUND_FLOOR) * step
        maximum = (maximum / step).to_integral_value(rounding=ROUND_CEILING) * step
        missing = max(Decimal(0), AXIS_MIN_INTERVALS - (maximum - minimum) / step)
        if high > 0:
            maximum += missing * step
        else:
            minimum -= missing * step
        return minimum, maximum, step


def _configure_template(presentation: PresentationDocument) -> None:
    original_width = presentation.slide_width
    presentation.slide_width = WIDTH
    presentation.slide_height = Inches(7.5)
    presentation._element.sldSz.set('type', 'custom')
    for layout in tuple(presentation.slide_masters) + tuple(presentation.slide_layouts):
        for element in layout._element.xpath('.//a:xfrm/a:off'):
            element.set('x', str(int(element.get('x')) * WIDTH // original_width))
        for element in layout._element.xpath('.//a:xfrm/a:ext'):
            element.set('cx', str(int(element.get('cx')) * WIDTH // original_width))
    for master in presentation.slide_masters:
        theme_part = master.part.part_related_by(RELATIONSHIP_TYPE.THEME)
        theme = parse_xml(theme_part.blob)
        for font in theme.xpath('./a:themeElements/a:fontScheme/*/a:latin'):
            font.set('typeface', FONT_NAME)
        theme_part.blob = serialize_part_xml(theme)


def _get_number_format(values: Sequence[Decimal]) -> str:
    with localcontext() as context:
        context.prec = 28
        places = max(max(0, -value.normalize().as_tuple().exponent) for value in values)
    return '0' if places == 0 else '0.' + '0' * places


def _get_numeric_text_width(text: str) -> Decimal:
    digits = sum(char.isdigit() for char in text)
    signs = text.count('-')
    separators = text.count('.')
    width = digits * AXIS_DIGIT_WIDTH + signs * AXIS_SIGN_WIDTH + separators * AXIS_SEPARATOR_WIDTH
    return width * AXIS_FONT_SIZE / POINTS_PER_INCH


def _get_fixed_plot_layout(
        minimum: Decimal, maximum: Decimal, step: Decimal) -> Tuple[Decimal, Decimal, Decimal, Decimal, int]:
    with localcontext() as context:
        context.prec = 28
        chart_width = Decimal(WIDTH) / Inches(1) - CONTENT_MARGINS - CHART_GUTTER
        number_format = _get_number_format((minimum, maximum, step))
        places = len(number_format.partition('.')[2])
        lower_text = format(minimum, f'.{places}f')
        upper_text = format(maximum, f'.{places}f')
        left = max(AXIS_MIN_MARGIN, _get_numeric_text_width(lower_text) / 2 + AXIS_TICK_MARGIN)
        right = max(AXIS_MIN_MARGIN, _get_numeric_text_width(upper_text) / 2 + AXIS_TICK_MARGIN)
        if minimum < 0:
            left = max(left, _get_numeric_text_width(format(minimum, '.2f')) + AXIS_VALUE_MARGIN)
        if maximum > 0:
            right = max(right, _get_numeric_text_width(format(maximum, '.2f')) + AXIS_VALUE_MARGIN)
        width = chart_width - left - right
        intervals = int((maximum - minimum) / step)
        tick_width = max(_get_numeric_text_width(format(minimum + index * step, f'.{places}f'))
                         for index in range(intervals + 1))
        rotation = AXIS_LABEL_ROTATION if tick_width + AXIS_VALUE_MARGIN > width / intervals else 0
        bottom = CHART_HEIGHT * AXIS_BOTTOM_MARGIN
        if rotation:
            diagonal = Decimal(2).sqrt() / 2
            label_height = (tick_width + AXIS_FONT_SIZE / POINTS_PER_INCH) * diagonal
            bottom = max(bottom, AXIS_LEGEND_MARGIN + label_height)
        height = CHART_HEIGHT * (1 - AXIS_TOP_MARGIN) - bottom
        return left / chart_width, AXIS_TOP_MARGIN, width / chart_width, height / CHART_HEIGHT, rotation


def _add_category_labels(
        slide: Slide, page: Sequence[Metric],
        layout: Optional[Tuple[Decimal, Decimal, Decimal, Decimal, int]]) -> None:
    with localcontext() as context:
        context.prec = 28
        for row, metric in enumerate(page):
            center = 1.15 + 5.1 * (.03 + .80 * (row + .5) / len(page))
            height = 1.0
            if layout is not None:
                row_height = CHART_HEIGHT * layout[3] / len(page)
                center = CHART_TOP + CHART_HEIGHT * layout[1] + row_height * (Decimal(row) + Decimal('0.5'))
                height = min(Decimal(1), row_height)
            frame = _text(slide, metric.name, .5, center - height / 2, 3.9, height)
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            if layout is not None:
                for paragraph in frame.paragraphs:
                    paragraph.line_spacing = Pt(CATEGORY_LINE_HEIGHT)


def _configure_fixed_chart(
        chart: Chart, bounds: Tuple[Decimal, Decimal], step: Decimal,
        layout: Tuple[Decimal, Decimal, Decimal, Decimal, int]) -> None:
    chart.value_axis.tick_labels.number_format = _get_number_format((bounds[0], bounds[1], step))
    chart.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.HIGH
    properties = chart.value_axis._element.get_or_add_txPr()
    properties.bodyPr.set('rot', str(layout[4]))
    _plot_layout(chart, layout[:4])


def _chart_slide(
        prs: PresentationDocument, unit: str, page: Sequence[Metric], period: str,
        number: int, index: int, count: int, axis_range: Optional[Tuple[Decimal, Decimal]] = None) -> None:
    slide = _new_slide(prs, f'План и факт · {unit} · график {index} из {count}', period, number)
    values = tuple(value for metric in page for value in (metric.plan, metric.actual))
    all_zero = all(value == 0 for value in values)
    number_format = '0' if all_zero else '0.00'
    layout = None
    if axis_range is not None:
        number_format = _get_number_format(values)
        step = report_data.calculate_fixed_axis_step(axis_range[0], axis_range[1])
        layout = _get_fixed_plot_layout(axis_range[0], axis_range[1], step)
    chart_data = _SafeChartData(number_format=number_format)
    chart_data.categories = [metric.name for metric in page]
    chart_data.add_series('План', [float(metric.plan) for metric in page])
    chart_data.add_series('Факт', [float(metric.actual) for metric in page])
    _add_category_labels(slide, page, layout)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(4.55), Inches(1.15),
                                   Inches(CONTENT_WIDTH - 4.05), Inches(5.1), chart_data).chart
    _font(chart.font, 14)
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    _font(chart.legend.font, 14)
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.overlap = 0
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.show_value = True
    labels.show_category_name = False
    labels.show_series_name = False
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    _font(labels.font, 14)
    for series, color in zip(chart.series, (PLAN, ACTUAL)):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
        series.format.line.fill.background()
        series.invert_if_negative = False
    category = chart.category_axis
    category.reverse_order = True
    # Текстовые блоки переносят длинные названия без автоматического уменьшения шрифта.
    category.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    category.major_tick_mark = category.minor_tick_mark = XL_TICK_MARK.NONE
    _font(category.tick_labels.font, 14)
    axis = chart.value_axis
    axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    axis.major_tick_mark = axis.minor_tick_mark = XL_TICK_MARK.NONE
    axis.tick_labels.number_format = number_format
    _font(axis.tick_labels.font, 14)
    if axis_range is None:
        scale = _calculate_axis_scale(values)
    else:
        scale = (axis_range[0], axis_range[1], step)
    axis.minimum_scale = float(scale[0])
    axis.maximum_scale = float(scale[1])
    axis.major_unit = float(scale[2])
    if layout is None:
        _plot_layout(chart)
    else:
        _configure_fixed_chart(chart, axis_range, step, layout)
        note = (f'Шкала фиксирована: {report_data.format_number(axis_range[0])}…'
                f'{report_data.format_number(axis_range[1])} {unit}')
        _text(slide, note, .5, AXIS_NOTE_TOP, CONTENT_WIDTH, AXIS_NOTE_HEIGHT)


def build_presentation(
        metrics: Sequence[Metric], title: str, period: str, comment: str,
        axis_ranges: Optional[Mapping[str, Tuple[Decimal, Decimal]]] = None) -> bytes:
    """Build a widescreen presentation with editable tables and native charts

    Args:
        metrics: Validated report metrics
        title: Required report title
        period: Required reporting period
        comment: Optional title-slide paragraphs
        axis_ranges: Optional fixed bounds for normalized current units

    Returns:
        Complete PPTX bytes

    Raises:
        ValueError: If report fields, metric count, fixed ranges or data bounds are invalid
    """
    fields = validate_report(title, period, comment)
    title = fields[0]
    period = fields[1]
    comment = fields[2]
    if not 1 <= len(metrics) <= 12:
        raise ValueError('Для презентации нужны от 1 до 12 показателей.')
    report_data.validate_axis_ranges(metrics, axis_ranges)
    prs = Presentation()
    _configure_template(prs)
    timestamp = datetime.now(timezone.utc)
    prs.core_properties.title = title
    prs.core_properties.author = PRODUCT_NAME
    prs.core_properties.last_modified_by = PRODUCT_NAME
    prs.core_properties.comments = PRODUCT_NAME
    prs.core_properties.created = timestamp
    prs.core_properties.modified = timestamp
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    _text(slide, title, .5, .6, CONTENT_WIDTH, 2.1, 32, True)
    _text(slide, period, .5, 2.9, CONTENT_WIDTH, .8, 22)
    if comment:
        # Пустые абзацы не вытесняют содержательный текст за границу титула.
        paragraphs = '\n'.join(line for line in comment.splitlines() if line.strip())
        _text(slide, paragraphs, .5, 3.85, CONTENT_WIDTH, 3.15)
    pages = slide_pages(metrics)
    tables = pages[0]
    charts = pages[1]
    for index, page in enumerate(tables, 1):
        _table_slide(prs, page, period, len(prs.slides) + 1, index, len(tables))
    for index, (unit, page) in enumerate(charts, 1):
        bounds = axis_ranges.get(unit.casefold()) if axis_ranges else None
        _chart_slide(prs, unit, page, period, len(prs.slides) + 1, index, len(charts), bounds)
    properties_part = prs.part.package.part_related_by(RELATIONSHIP_TYPE.EXTENDED_PROPERTIES)
    properties = parse_xml(properties_part.blob)
    properties.find(PROPERTIES_NAMESPACE + 'PresentationFormat').text = 'On-screen Show (16:9)'
    properties.find(PROPERTIES_NAMESPACE + 'Slides').text = str(len(prs.slides))
    properties.find(PROPERTIES_NAMESPACE + 'Application').text = PRODUCT_NAME
    properties_part.blob = serialize_part_xml(properties)
    output = BytesIO()
    prs.save(output)
    return output.getvalue()
