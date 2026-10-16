from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION
from pptx.util import Inches, Pt

from data import Metric, parse_csv
from slides import build_presentation, slide_pages


SAMPLE = Path(__file__).resolve().parents[1] / 'samples/weekly-report.csv'
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'ep': 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties',
      's': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def metrics(count, same_unit=False):
    return tuple(Metric(f'Показатель {n + 1}', Decimal('1.25'), Decimal('-2.5'),
                        'шт.' if same_unit else f'ед. {n + 1}') for n in range(count))


def objects(prs):
    tables, charts = [], []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table:
                tables.append(shape.table)
            if shape.has_chart:
                charts.append(shape.chart)
    return tables, charts


def workbook_cells(blob):
    with ZipFile(BytesIO(blob)) as archive:
        strings = [''.join(node.itertext()) for node in
                   ET.fromstring(archive.read('xl/sharedStrings.xml')).findall('s:si', NS)]
        root = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        cells = {}
        for cell in root.findall('.//s:c', NS):
            value = cell.find('s:v', NS).text
            cells[cell.attrib['r']] = strings[int(value)] if cell.attrib.get('t') == 's' else Decimal(value)
        return cells


class PresentationTests(unittest.TestCase):
    def test_axis_ticks_are_aligned_and_readable_at_all_numeric_limits(self) -> None:
        """Keep zero on a readable tick without duplicate cent-scale labels"""
        # ARRANGE
        cases = (('120', '135'), ('-10', '5'), ('-10', '-7'), ('0.01', '0'),
                 ('-0.01', '0'), ('-0.01', '0.01'), ('-1000000000', '999999999.99'))
        for plan, actual in cases:
            source = (Metric('Показатель', Decimal(plan), Decimal(actual), 'шт.'),)
            # ACT
            prs = Presentation(BytesIO(build_presentation(source, 'Отчёт', 'Неделя', '')))
            chart = next(shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart)
            axis = chart.value_axis
            minimum = Decimal(str(axis.minimum_scale))
            maximum = Decimal(str(axis.maximum_scale))
            step = Decimal(str(axis.major_unit))
            intervals = (maximum - minimum) / step
            labels = [format(minimum + index * step, '.2f') for index in range(int(intervals) + 1)]
            # ASSERT
            with self.subTest(plan=plan, actual=actual):
                self.assertGreaterEqual(step, Decimal('0.01'))
                self.assertIn(step / Decimal(10) ** step.adjusted(), (Decimal(1), Decimal(2), Decimal(5)))
                self.assertEqual(minimum % step, 0)
                self.assertEqual(maximum % step, 0)
                self.assertLessEqual(minimum, min(Decimal(plan), Decimal(actual), Decimal(0)))
                self.assertGreaterEqual(maximum, max(Decimal(plan), Decimal(actual), Decimal(0)))
                self.assertGreaterEqual(intervals, 3)
                self.assertLessEqual(intervals, 5)
                self.assertEqual(len(set(labels)), len(labels))
                if max(Decimal(plan), Decimal(actual)) <= 0:
                    self.assertEqual(maximum, 0)
                if min(Decimal(plan), Decimal(actual)) >= 0:
                    self.assertEqual(minimum, 0)

    def test_added_slides_inherit_widescreen_layout_and_arial_theme(self) -> None:
        """Keep new PowerPoint slides consistent with the generated deck"""
        # ARRANGE
        blob = build_presentation(metrics(1), 'Отчёт', 'Неделя', '')
        # ACT
        prs = Presentation(BytesIO(blob))
        layout = next(item for item in prs.slide_layouts if item.name == 'Title and Content')
        slide = prs.slides.add_slide(layout)
        # ASSERT
        self.assertEqual([(shape.left, shape.width) for shape in slide.placeholders],
                         [(609600, 10972800), (609600, 10972800)])
        self.assertEqual(prs._element.sldSz.get('type'), 'custom')
        with ZipFile(BytesIO(blob)) as archive:
            theme = ET.fromstring(archive.read('ppt/theme/theme1.xml'))
            self.assertEqual([font.get('typeface') for font in
                              theme.findall('./a:themeElements/a:fontScheme/*/a:latin', NS)],
                             ['Arial', 'Arial'])

    def test_metadata_describes_this_export_instead_of_the_library_template(self) -> None:
        """Store product identity and current UTC timestamps in a real PPTX"""
        # ARRANGE
        timestamp = datetime(2026, 10, 14, 12, 30, tzinfo=timezone.utc)
        # ACT
        with patch('slides.datetime') as clock:
            clock.now.return_value = timestamp
            blob = build_presentation(metrics(1), 'Отчёт', 'Неделя', '')
        properties = Presentation(BytesIO(blob)).core_properties
        # ASSERT
        clock.now.assert_called_once_with(timezone.utc)
        self.assertEqual(properties.last_modified_by, 'Report Slides')
        self.assertEqual(properties.author, 'Report Slides')
        self.assertEqual(properties.comments, 'Report Slides')
        self.assertEqual(properties.created, timestamp.replace(tzinfo=None))
        self.assertEqual(properties.modified, timestamp.replace(tzinfo=None))
        with ZipFile(BytesIO(blob)) as archive:
            extended = ET.fromstring(archive.read('docProps/app.xml'))
            self.assertEqual(extended.find('ep:PresentationFormat', NS).text, 'On-screen Show (16:9)')
            self.assertEqual(extended.find('ep:Slides', NS).text, '3')
            self.assertEqual(extended.find('ep:Application', NS).text, 'Report Slides')

    def test_sample_native_objects_numbers_and_five_slides(self):
        source = parse_csv(SAMPLE.read_bytes(), SAMPLE.name, ';')
        blob = build_presentation(source, 'Отчёт', '5–11 октября 2026', 'Комментарий')
        prs = Presentation(BytesIO(blob))
        self.assertEqual(len(prs.slides), 5)
        self.assertEqual((prs.slide_width, prs.slide_height), (12192000, Inches(7.5)))
        tables, charts = objects(prs)
        self.assertEqual(len(tables), 1)
        self.assertEqual([[cell.text for cell in row.cells] for row in tables[0].rows], [
            ['Показатель', 'Единица', 'План', 'Факт', 'Отклонение', 'Факт/план, %'],
            ['Заявки', 'шт.', '120', '135', '15', '112,5'],
            ['Выручка', 'руб.', '250000,5', '275000,75', '25000,25', '110,0'],
            ['Время ответа', 'ч', '2,5', '2', '-0,5', '80,0'],
            ['Возвраты', 'шт.', '0', '3', '3', '—'],
        ])
        self.assertEqual(len(charts), 3)
        expected = [(['Заявки', 'Возвраты'], [(120, 0), (135, 3)]),
                    (['Выручка'], [(250000.5,), (275000.75,)]),
                    (['Время ответа'], [(2.5,), (2,)])]
        for chart, (names, values) in zip(charts, expected):
            self.assertEqual(chart.chart_type, XL_CHART_TYPE.BAR_CLUSTERED)
            self.assertEqual([category.label for category in chart.plots[0].categories], names)
            self.assertEqual([series.name for series in chart.series], ['План', 'Факт'])
            self.assertEqual([series.values for series in chart.series], values)
        for index, slide in enumerate(prs.slides, 1):
            text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
            if index > 1:
                self.assertIn('5–11 октября 2026', text)
                self.assertIn(str(index), text)
        self.assertIn('Комментарий', '\n'.join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame))

    def test_twelve_rows_pagination_preserves_all_names_and_group_order(self):
        for same_unit, count in ((True, 7), (False, 16)):
            with self.subTest(same_unit=same_unit):
                source = metrics(12, same_unit)
                tables, pages = slide_pages(source)
                self.assertEqual(len(tables), 3)
                self.assertEqual(len(pages), 3 if same_unit else 12)
                prs = Presentation(BytesIO(build_presentation(source, 'Отчёт', 'Неделя', '')))
                self.assertEqual(len(prs.slides), count)
                native_tables, charts = objects(prs)
                self.assertEqual([row.cells[0].text for table in native_tables for row in list(table.rows)[1:]],
                                 [f'Показатель {n + 1}' for n in range(12)])
                self.assertEqual([cat.label for chart in charts for cat in chart.plots[0].categories],
                                 [f'Показатель {n + 1}' for n in range(12)])
                self.assertTrue(all(len(chart.plots[0].categories) <= 4 for chart in charts))

    def test_chart_cache_and_workbook_numbers_match_csv_to_two_places(self):
        source = (Metric('Крайний', Decimal('-1000000000'), Decimal('999999999.99'), '%'),
                  Metric('Малый', Decimal('0.01'), Decimal('0.29'), '%'))
        blob = build_presentation(source, 'Отчёт', 'Неделя', '')
        with ZipFile(BytesIO(blob)) as archive:
            root = ET.fromstring(archive.read('ppt/charts/chart1.xml'))
            cache = [[Decimal(node.text).quantize(Decimal('0.01')) for node in
                      series.findall('c:val/c:numRef/c:numCache/c:pt/c:v', NS)]
                     for series in root.findall('.//c:ser', NS)]
            self.assertEqual(cache, [[Decimal('-1000000000.00'), Decimal('0.01')],
                                     [Decimal('999999999.99'), Decimal('0.29')]])
            xlsx_name = next(name for name in archive.namelist() if name.endswith('.xlsx'))
            cells = workbook_cells(archive.read(xlsx_name))
            self.assertEqual([cells[k].quantize(Decimal('0.01')) for k in ('B2', 'B3', 'C2', 'C3')],
                             [Decimal('-1000000000.00'), Decimal('0.01'),
                              Decimal('999999999.99'), Decimal('0.29')])
            self.assertEqual([cells[k] for k in ('B1', 'C1', 'A2', 'A3')], ['План', 'Факт', 'Крайний', 'Малый'])
            for node in root.findall('.//c:numFmt', NS):
                self.assertNotIn('%', node.attrib['formatCode'])
                self.assertNotIn('E', node.attrib['formatCode'].upper())

    def test_literal_text_in_embedded_books_no_formulas_links_or_temp_files(self):
        names = ['=1+1', 'https://example.test', '<А & "Б">', '+SUM(A1)']
        source = tuple(Metric(name, Decimal('25'), Decimal('30'), '%') for name in names)
        with patch('tempfile.mkstemp', side_effect=AssertionError('Data written to disk')):
            blob = build_presentation(source, '=Отчёт & <текст>', 'Неделя', '<a href="x">Текст</a>')
        with ZipFile(BytesIO(blob)) as archive:
            for name in archive.namelist():
                if name.endswith('.rels'):
                    self.assertNotIn(b'TargetMode="External"', archive.read(name))
                if name.endswith('.xlsx'):
                    cells = workbook_cells(archive.read(name))
                    self.assertEqual([cells[f'A{n + 2}'] for n in range(4)], names)
                    self.assertEqual(cells['B2'], Decimal('25'))
                    with ZipFile(BytesIO(archive.read(name))) as book:
                        for member in book.namelist():
                            if member.startswith('xl/worksheets/') and member.endswith('.xml'):
                                root = ET.fromstring(book.read(member))
                                self.assertFalse(root.findall('.//s:f', NS))
                                self.assertFalse(root.findall('.//s:hyperlink', NS))
                        self.assertFalse(any('externalLink' in member or 'vbaProject' in member for member in book.namelist()))
            self.assertFalse(any('vbaProject' in name for name in archive.namelist()))

    def test_axes_include_zero_and_all_zero_chart_has_zero_labels(self):
        for values in (('0', '0'), ('-5', '-2'), ('-5', '2'), ('1', '2')):
            with self.subTest(values=values):
                source = (Metric('А', Decimal(values[0]), Decimal(values[1]), 'шт.'),)
                prs = Presentation(BytesIO(build_presentation(source, 'Отчёт', 'Неделя', '')))
                chart = objects(prs)[1][0]
                self.assertLessEqual(chart.value_axis.minimum_scale, 0)
                self.assertGreaterEqual(chart.value_axis.maximum_scale, 0)
                self.assertTrue(chart.plots[0].has_data_labels)
                self.assertTrue(chart.category_axis.reverse_order)
                with ZipFile(BytesIO(build_presentation(source, 'Отчёт', 'Неделя', ''))) as archive:
                    xml = ET.fromstring(archive.read('ppt/charts/chart1.xml'))
                    self.assertEqual(xml.find('.//c:catAx/c:crosses', NS).attrib['val'], 'autoZero')
                if values == ('0', '0'):
                    self.assertEqual((chart.value_axis.minimum_scale, chart.value_axis.maximum_scale), (0, 1))
                    self.assertEqual(chart.plots[0].data_labels.number_format, '0')

    def test_chart_names_have_room_to_wrap_and_native_data_keep_full_text(self):
        source = metrics(4, True)
        source = tuple(Metric(str(n + 1) + 'Ш' * 59, m.plan, m.actual, m.unit) for n, m in enumerate(source))
        prs = Presentation(BytesIO(build_presentation(source, 'Отчёт', 'Неделя', '')))
        slide = prs.slides[-1]
        labels = [shape for shape in slide.shapes if shape.has_text_frame and shape.text in {m.name for m in source}]
        self.assertEqual([label.text for label in labels], [m.name for m in source])
        self.assertTrue(all(label.width >= Inches(3.5) and label.height >= Inches(.9) for label in labels))
        chart_shape = next(shape for shape in slide.shapes if shape.has_chart)
        self.assertGreaterEqual(chart_shape.left, Inches(4.5))
        self.assertEqual([cat.label for cat in chart_shape.chart.plots[0].categories], [m.name for m in source])

    def test_long_text_font_floors_margins_and_repeated_layout(self):
        source = tuple(Metric(f'{n + 1}' + 'Ш' * 59, Decimal('-999999999.99'),
                              Decimal('999999999.99'), 'Щ' * 16) for n in range(12))
        title, period, comment = 'Ш' * 100, 'Щ' * 60, '\n'.join(['Ш' * 65] * 6)
        blobs = [build_presentation(source, title, period, comment) for _ in range(2)]
        prs = Presentation(BytesIO(blobs[0]))
        other = Presentation(BytesIO(blobs[1]))
        self.assertEqual(len(prs.slides), 7)
        for slide, second in zip(prs.slides, other.slides):
            self.assertEqual([(s.left, s.top, s.width, s.height) for s in slide.shapes],
                             [(s.left, s.top, s.width, s.height) for s in second.shapes])
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, Inches(.4))
                self.assertGreaterEqual(shape.top, Inches(.4))
                self.assertLessEqual(shape.left + shape.width, prs.slide_width - Inches(.4))
                self.assertLessEqual(shape.top + shape.height, prs.slide_height - Inches(.4))
                frames = ([shape.text_frame] if shape.has_text_frame else [])
                if shape.has_table:
                    frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
                for frame in frames:
                    for paragraph in frame.paragraphs:
                        self.assertGreaterEqual(paragraph.font.size, Pt(16))
                        self.assertEqual(paragraph.font.name, 'Arial')
        self.assertEqual(objects(prs)[0][0].cell(1, 0).text, '1' + 'Ш' * 59)

    def test_export_validates_required_fields_and_nonempty_data(self):
        for source, title, period in (((), 'А', 'Б'), (metrics(1), '', 'Б'), (metrics(1), 'А', ''),
                                      (metrics(13), 'А', 'Б')):
            with self.subTest(title=title, period=period), self.assertRaises(ValueError):
                build_presentation(source, title, period, '')

    def test_blank_comment_paragraphs_do_not_push_meaningful_text_off_slide(self):
        comment = 'Начало' + '\n' * 30 + 'Конец комментария'
        prs = Presentation(BytesIO(build_presentation(metrics(1), 'Отчёт', 'Неделя', comment)))
        frame = next(shape.text_frame for shape in prs.slides[0].shapes
                     if shape.has_text_frame and 'Начало' in shape.text)
        self.assertEqual([p.text for p in frame.paragraphs], ['Начало', 'Конец комментария'])
        self.assertTrue(all(p.font.size >= Pt(16) for p in frame.paragraphs))


    def test_tab_heavy_comment_preserves_text_and_bounds_tab_width(self):
        comment = 'Начало' + '\t' * 389 + 'КОНЕЦ'
        prs = Presentation(BytesIO(build_presentation(metrics(1), 'Отчёт', 'Неделя', comment)))
        frame = next(shape.text_frame for shape in prs.slides[0].shapes
                     if shape.has_text_frame and shape.text.startswith('Начало'))
        self.assertEqual(frame.text, comment)
        paragraph = frame.paragraphs[0]
        tab_size = paragraph._p.pPr.get('defTabSz')
        self.assertIsNotNone(tab_size, 'Default half-inch tabs can push the final word off the title slide')
        self.assertLessEqual(int(tab_size), Pt(16))
        self.assertEqual(paragraph.font.size, Pt(16))


class FixedAxisTests(unittest.TestCase):
    """Verify fixed ranges in actual native presentations"""

    def test_same_ranges_preserve_geometry_across_weeks_and_pages(self) -> None:
        """Use identical axes and plot geometry with changed values and category counts"""
        # ARRANGE
        first = (Metric('А', Decimal(100), Decimal(135), 'шт.'),
                 Metric('Б', Decimal(0), Decimal(200), 'ШТ.'),
                 Metric('В', Decimal(10), Decimal(20), 'шт.'),
                 Metric('Г', Decimal(30), Decimal(40), 'шт.'),
                 Metric('Д', Decimal(50), Decimal(60), 'шт.'))
        second = (Metric('А', Decimal(100), Decimal(135), 'шт.'),
                  Metric('Б', Decimal(20), Decimal(150), 'шт.'))
        bounds = {'шт.': (Decimal(0), Decimal(200))}
        signatures = []
        # ACT
        for source in (first, second):
            blob = build_presentation(source, 'Отчёт', 'Неделя', '', axis_ranges=bounds)
            presentation = Presentation(BytesIO(blob))
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if shape.has_chart:
                        axis = shape.chart.value_axis
                        manual = shape.chart._chartSpace.chart.plotArea.find('c:layout', NS)
                        signatures.append((axis.minimum_scale, axis.maximum_scale, axis.major_unit,
                                           axis.tick_labels.number_format, ET.tostring(manual)))
        # ASSERT
        self.assertEqual(len(signatures), 3)
        self.assertEqual(len(set(signatures)), 1)
        self.assertEqual(bounds, {'шт.': (Decimal(0), Decimal(200))})

    def test_zero_fractional_and_mixed_unit_ranges(self) -> None:
        """Keep fractional ticks and exact fractional data on an integer-scaled axis"""
        # ARRANGE
        source = (Metric('Нули', Decimal(0), Decimal(0), 'шт.'),
                  Metric('Дробь', Decimal('120.25'), Decimal(135), 'руб.'),
                  Metric('Авто', Decimal(1), Decimal(2), 'ч'))
        ranges = {'шт.': (Decimal(0), Decimal('0.02')), 'руб.': (Decimal(0), Decimal(200))}
        signatures = []
        # ACT
        blob = build_presentation(source, 'Отчёт', 'Неделя', '', ranges)
        presentation = Presentation(BytesIO(blob))
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_chart:
                    axis = shape.chart.value_axis
                    signatures.append((axis.minimum_scale, axis.maximum_scale, axis.major_unit,
                                       axis.tick_labels.number_format,
                                       shape.chart.plots[0].data_labels.number_format))
        # ASSERT
        self.assertEqual(signatures, [(0.0, 0.02, 0.01, '0.00', '0'),
                                      (0.0, 200.0, 50.0, '0', '0.00'),
                                      (0.0, 3.0, 1.0, '0.00', '0.00')])
        with ZipFile(BytesIO(blob)) as archive:
            books = [workbook_cells(archive.read(name)) for name in archive.namelist() if name.endswith('.xlsx')]
            self.assertTrue(any(Decimal('120.25') in book.values() for book in books))

    def test_indicator_footer_labels_and_conditional_rotation(self) -> None:
        """Keep the scale note out of the footer and category rectangles within their rows"""
        # ARRANGE
        cases = ((Decimal(200), Decimal(50), 0), (Decimal(900000000), Decimal(100000000), -2700000))
        # ACT
        for maximum, step, rotation in cases:
            source = (Metric('А', Decimal(0), maximum, 'шт.'),
                      Metric('Б', Decimal(0), maximum, 'шт.'),
                      Metric('В', Decimal(0), maximum, 'шт.'),
                      Metric('Г', Decimal(0), maximum, 'шт.'))
            blob = build_presentation(source, 'Отчёт', 'Неделя', '', {'шт.': (Decimal(0), maximum)})
            presentation = Presentation(BytesIO(blob))
            # ASSERT
            for slide in presentation.slides:
                charts = [shape.chart for shape in slide.shapes if shape.has_chart]
                if not charts:
                    continue
                notes = [shape for shape in slide.shapes if shape.has_text_frame
                         and shape.text.startswith('Шкала фиксирована:')]
                self.assertEqual([(shape.top, shape.height) for shape in notes], [(Inches(6.34), Inches(.27))])
                for note in notes:
                    self.assertLessEqual(note.top + note.height, Inches(6.64))
                labels = [shape for shape in slide.shapes if shape.has_text_frame
                          and shape.text in ('А', 'Б', 'В', 'Г')]
                previous_bottom = 0
                for label in labels:
                    self.assertGreaterEqual(label.top, previous_bottom)
                    self.assertEqual([paragraph.line_spacing for paragraph in label.text_frame.paragraphs], [Pt(15)])
                    previous_bottom = label.top + label.height
                for chart in charts:
                    self.assertEqual(chart.value_axis.major_unit, float(step))
                    self.assertEqual(chart.value_axis._element.txPr.bodyPr.get('rot'), str(rotation))
                    self.assertEqual(chart.value_axis.tick_label_position, XL_TICK_LABEL_POSITION.HIGH)
                    manual = chart._chartSpace.chart.plotArea.find('c:layout/c:manualLayout', NS)
                    self.assertEqual([(tag, manual.find('c:' + tag, NS).get('val'))
                                      for tag in ('xMode', 'yMode', 'wMode', 'hMode')],
                                     [('xMode', 'edge'), ('yMode', 'edge'), ('wMode', 'factor'), ('hMode', 'factor')])

    def test_direct_failures_precede_creation_and_keep_report_error_order(self) -> None:
        """Validate metadata before ranges and reject clipped exports before constructing a package"""
        # ARRANGE
        source = (Metric('А', Decimal(201), Decimal(230), 'шт.'),)
        # ACT
        with patch('slides.Presentation') as constructor:
            with self.assertRaisesRegex(ValueError, 'Название'):
                build_presentation(source, '', 'Неделя', '', {'шт.': (Decimal(0), Decimal(117))})
            with self.assertRaisesRegex(ValueError, 'План: 201.*\n.*Факт: 230'):
                build_presentation(source, 'Отчёт', 'Неделя', '', {'шт.': (Decimal(0), Decimal(200))})
            # ASSERT
            constructor.assert_not_called()


if __name__ == '__main__':
    unittest.main()
