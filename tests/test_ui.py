from io import BytesIO
from pathlib import Path
import socket
import unittest
from unittest.mock import patch

from pptx import Presentation
from streamlit.testing.v1 import AppTest

import data as report_data


ROOT = Path(__file__).resolve().parents[1]
SAMPLE = (ROOT / 'samples/weekly-report.csv').read_bytes()


class InterfaceTests(unittest.TestCase):
    def test_field_feedback_waits_for_editing_and_uses_an_error_alert(self) -> None:
        """Avoid premature field errors while retaining feedback after user interaction"""
        # ARRANGE
        app = self.app()
        self.assertFalse(app.error)
        self.assertFalse(any('Период: заполните поле' in message.value for message in app.caption))
        self.upload(app)
        self.assertFalse(app.error)
        # ACT
        app.text_input(key='period').set_value('Неделя').run()
        app.text_input(key='period').set_value('').run()
        # ASSERT
        self.assertEqual([message.value for message in app.error], ['Период: заполните поле.'])
        self.assertTrue(app.button(key='generate').disabled)
        app.button(key='new_csv').click().run()
        self.assertFalse(app.error)
        self.assertFalse(app.session_state['edited_fields'])

    def test_feedback_does_not_mark_other_untouched_fields_as_invalid(self) -> None:
        """Validate edited fields independently even when the period is still empty"""
        # ARRANGE
        app = self.upload(self.app())
        # ACT
        app.text_input(key='title').set_value('Другой отчёт').run()
        # ASSERT
        self.assertFalse(app.error)
        self.assertTrue(app.button(key='generate').disabled)
        app.text_area(key='comment').set_value('а' * 401).run()
        self.assertEqual([message.value for message in app.error],
                         ['Комментарий: допускается не больше 400 символов.'])

    def test_title_feedback_survives_upload_and_reset(self) -> None:
        """Keep explaining an invalid title while its value survives CSV replacement"""
        # ARRANGE
        app = self.app()
        title_error = ['Название отчёта: заполните поле.']
        # ACT
        app.text_input(key='title').set_value('').run()
        # ASSERT
        self.assertEqual([message.value for message in app.error], title_error)
        self.upload(app)
        app.text_input(key='period').set_value('Неделя').run()
        self.assertEqual([message.value for message in app.error], title_error)
        self.assertTrue(app.button(key='generate').disabled)
        app.button(key='new_csv').click().run()
        self.assertEqual([message.value for message in app.error], title_error)
        self.upload(app)
        app.text_input(key='period').set_value('Неделя').run()
        self.assertEqual([message.value for message in app.error], title_error)
        app.text_input(key='title').set_value('Отчёт').run()
        self.assertFalse(app.error)
        self.assertFalse(app.button(key='generate').disabled)

    def test_generation_logs_a_traceback_without_user_content(self) -> None:
        """Retain a diagnosable stack while hiding sensitive exception text"""
        # ARRANGE
        app = self.generated()
        previous = app.session_state['pptx']
        # ACT
        with self.assertLogs('report_slides', level='ERROR') as messages:
            with patch('slides.build_presentation', side_effect=RuntimeError('SECRET XML <csv>')):
                app.button(key='generate').click().run()
        # ASSERT
        self.assertIn('Traceback', '\n'.join(messages.output))
        self.assertNotIn('SECRET', '\n'.join(messages.output))
        self.assertTrue(all(record.exc_info for record in messages.records))
        self.assertTrue(all('SECRET' not in message.value for message in app.error))
        self.assertEqual(app.session_state['pptx'], previous)

    def app(self):
        app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=15).run()
        self.assertFalse(app.exception)
        return app

    def upload(self, app, raw=SAMPLE, filename='weekly.csv'):
        if not app.file_uploader:
            app.button(key='new_csv').click().run()
        app.file_uploader[0].set_value((filename, raw, 'text/csv')).run()
        self.assertFalse(app.exception)
        return app

    def ready(self):
        app = self.upload(self.app())
        app.text_input(key='period').set_value('5–11 октября 2026').run()
        app.text_area(key='comment').set_value('Комментарий').run()
        self.assertFalse(app.exception)
        return app

    def generated(self):
        app = self.ready()
        app.button(key='generate').click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(Presentation(BytesIO(app.session_state['pptx'])).slides), 5)
        return app

    def assert_no_export(self, app):
        self.assertIsNone(app.session_state['pptx'])
        self.assertFalse(any(button.label == 'Скачать PPTX' for button in app.download_button))

    def test_empty_app_and_session_isolation(self):
        app = self.app()
        self.assertEqual(app.title[0].value, 'Report Slides')
        self.assertEqual(app.text_input(key='title').value, 'Еженедельный отчёт')
        self.assertEqual(app.text_input(key='period').value, '')
        self.assertTrue(app.button(key='generate').disabled)
        self.assertIsNone(app.session_state['metrics'])
        self.assert_no_export(app)
        self.assertTrue(any('сессия' in message.value for message in app.caption))
        another = self.app()
        self.upload(app)
        self.assertIsNone(another.session_state['metrics'])

    def test_metadata_fields_wait_for_valid_csv_during_replacement(self):
        app = self.app()
        self.assertTrue(app.text_input(key='period').disabled)
        self.assertTrue(app.text_area(key='comment').disabled)
        self.assertFalse(app.text_input(key='title').disabled)
        self.upload(app)
        self.assertFalse(app.text_input(key='period').disabled)
        self.assertFalse(app.text_area(key='comment').disabled)
        app.button(key='new_csv').click().run()
        self.assertTrue(app.text_input(key='period').disabled)
        self.upload(app, SAMPLE.replace(b'120', b'NaN'))
        self.assertTrue(app.text_input(key='period').disabled)
        self.assertTrue(app.text_area(key='comment').disabled)

    def test_real_upload_preview_composition_and_export_only_by_button(self):
        app = self.ready()
        self.assert_no_export(app)
        preview = app.dataframe[0].value
        self.assertEqual(preview['Отклонение'].tolist(), ['15', '25000,25', '-0,5', '3'])
        self.assertEqual(preview['Факт/план, %'].tolist(), ['112,5', '110,0', '80,0', '—'])
        self.assertTrue(any('5 слайдов' in text.value for text in app.text))
        app.button(key='generate').click().run()
        blob = app.session_state['pptx']
        self.assertEqual(len(Presentation(BytesIO(blob)).slides), 5)
        app.run()
        self.assertEqual(app.session_state['pptx'], blob)
        self.assertTrue(app.download_button(key='download_pptx').proto.url.endswith('.pptx'))
        self.assertTrue(app.download_button(key='download_pptx').proto.ignore_rerun)

    def test_replacing_filename_or_bytes_clears_fields_and_export_but_preserves_title(self):
        for raw, filename in ((SAMPLE, 'new.csv'), (SAMPLE.replace(b'120', b'121'), 'weekly.csv')):
            with self.subTest(filename=filename):
                app = self.generated()
                self.upload(app, raw, filename)
                self.assert_no_export(app)
                self.assertEqual(app.text_input(key='title').value, 'Еженедельный отчёт')
                self.assertEqual(app.text_input(key='period').value, '')
                self.assertEqual(app.text_area(key='comment').value, '')
                self.assertEqual(len(app.session_state['metrics']), 4)
                self.assertTrue(app.button(key='generate').disabled)

    def test_invalid_replacement_removes_old_data_and_download(self):
        app = self.generated()
        self.upload(app, SAMPLE.replace(b'120', b'NaN'))
        self.assertIsNone(app.session_state['metrics'])
        self.assertEqual(len(app.dataframe), 0)
        self.assert_no_export(app)
        self.assertTrue(app.error)
        self.assertIn('Строка 2', app.error[0].value)
        self.assertEqual(app.text_input(key='period').value, '')

    def test_removal_and_delimiter_change_clear_old_result(self):
        app = self.generated()
        app.button(key='new_csv').click().run()
        self.assertFalse(app.exception)
        self.assertIsNone(app.session_state['metrics'])
        self.assert_no_export(app)
        self.assertEqual(app.text_input(key='period').value, '')
        app = self.generated()
        app.segmented_control(key='delimiter').set_value(',').run()
        self.assertFalse(app.exception)
        self.assertIsNone(app.session_state['metrics'])
        self.assert_no_export(app)
        self.assertEqual(app.text_area(key='comment').value, 'Комментарий')
        self.assertEqual(app.text_input(key='period').value, '5–11 октября 2026')
        app.segmented_control(key='delimiter').set_value(';').run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.session_state['metrics']), 4)
        self.assertEqual(app.text_area(key='comment').value, 'Комментарий')
        self.assertEqual(app.text_input(key='period').value, '5–11 октября 2026')
        self.assert_no_export(app)

    def test_report_field_changes_invalidate_result_keep_data(self):
        for key, value in (('title', 'Другой отчёт'), ('period', 'Новая неделя'),
                            ('comment', 'Новый комментарий')):
            with self.subTest(key=key):
                app = self.generated()
                before = app.session_state['metrics']
                widget = app.text_area(key=key) if key == 'comment' else app.text_input(key=key)
                widget.set_value(value).run()
                self.assertFalse(app.exception)
                self.assert_no_export(app)
                self.assertEqual(app.session_state['metrics'], before)

    def test_invalid_report_fields_disable_generation_without_losing_data(self):
        for key, value in (('title', ' '), ('period', ''), ('comment', 'а' * 401)):
            with self.subTest(key=key):
                app = self.ready()
                widget = app.text_area(key=key) if key == 'comment' else app.text_input(key=key)
                widget.set_value(value).run()
                self.assertFalse(app.exception)
                self.assertTrue(app.button(key='generate').disabled)
                self.assertEqual(len(app.session_state['metrics']), 4)
                self.assert_no_export(app)

    def test_generation_failure_preserves_inputs_and_previous_success_only_if_unchanged(self):
        app = self.generated()
        previous = app.session_state['pptx']
        with patch('slides.build_presentation', side_effect=RuntimeError('SECRET XML <csv>')):
            app.button(key='generate').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.session_state['pptx'], previous)
            self.assertTrue(any(button.label == 'Скачать PPTX' for button in app.download_button))
            self.assertTrue(app.error)
            self.assertNotIn('SECRET', app.error[0].value)
            app.text_input(key='period').set_value('Новая неделя').run()
            app.button(key='generate').click().run()
            self.assert_no_export(app)
        self.assertEqual(app.text_input(key='period').value, 'Новая неделя')
        self.assertEqual(app.text_area(key='comment').value, 'Комментарий')
        self.assertEqual(len(app.session_state['metrics']), 4)

    def test_no_network_during_upload_preview_generation_and_rerun(self):
        app = self.app()
        with patch.object(socket.socket, 'connect', side_effect=AssertionError('External network')):
            self.upload(app)
            app.text_input(key='period').set_value('Неделя').run()
            app.button(key='generate').click().run()
            app.run()
        self.assertFalse(app.exception)
        self.assertEqual(len(Presentation(BytesIO(app.session_state['pptx'])).slides), 5)

    def test_preview_does_not_interpret_user_markdown_or_urls(self):
        app = self.upload(self.app(), 'показатель;план;факт;единица\n**А**;1;2;<b>ч</b>\nhttps://example.test;1;2;шт.\n'.encode())
        preview = app.dataframe[0].value
        self.assertEqual(preview['Показатель'].tolist(), ['**А**', 'https://example.test'])
        self.assertEqual(preview['Единица'].tolist(), ['<b>ч</b>', 'шт.'])
        self.assertEqual(len(app.get('link_button')), 0)


    def test_replacement_starts_by_clearing_report_before_new_file_is_selected(self):
        app = self.generated()
        self.assertEqual(len(app.file_uploader), 0)
        previous_revision = app.session_state['upload_revision']
        app.button(key='new_csv').click().run()
        self.assertFalse(app.exception)
        self.assertIsNone(app.session_state['metrics'])
        self.assertIsNone(app.session_state['source_file'])
        self.assert_no_export(app)
        self.assertIsNone(app.file_uploader[0].value)
        self.assertFalse(app.file_uploader[0].disabled)
        self.assertEqual(app.session_state['upload_revision'], previous_revision + 1)
        self.assertEqual(app.text_input(key='title').value, 'Еженедельный отчёт')
        self.assertEqual(app.text_input(key='period').value, '')
        self.assertEqual(app.text_area(key='comment').value, '')
        self.upload(app, SAMPLE.replace(b'120', b'NaN'))
        self.assertIsNone(app.session_state['metrics'])
        self.assert_no_export(app)


    def _set_fixed_axis(self, app: AppTest, unit: str = 'шт.', maximum: str = '200') -> AppTest:
        app.checkbox(key='axis_fixed_' + unit).set_value(True).run()
        app.text_input(key='axis_max_' + unit).set_value(maximum).run()
        self.assertFalse(app.exception)
        return app

    def test_fixed_controls_validation_expansion_and_real_export(self) -> None:
        """Open invalid controls and export the requested native range after correction"""
        # ARRANGE
        app = self.ready()
        # ACT
        app.checkbox(key='axis_fixed_шт.').set_value(True).run()
        # ASSERT
        self.assertTrue(app.button(key='generate').disabled)
        self.assertTrue(app.button(key='axis_save').disabled)
        self.assertTrue(any('Максимум' in message.value for message in app.error))
        self.assertEqual(len(app.error), 1)
        self.assertTrue(next(iter(app.expander)).proto.expanded)
        app.text_input(key='axis_max_шт.').set_value('200').run()
        self.assertFalse(app.error)
        app.button(key='generate').click().run()
        self.assertFalse(app.exception)
        charts = [shape.chart for slide in Presentation(BytesIO(app.session_state['pptx'])).slides
                  for shape in slide.shapes if shape.has_chart]
        chart = next(chart for chart in charts
                     if 'Заявки' in [category.label for category in chart.plots[0].categories])
        self.assertEqual((chart.value_axis.minimum_scale, chart.value_axis.maximum_scale, chart.value_axis.major_unit),
                         (0.0, 200.0, 50.0))

    def test_axis_edits_and_returning_values_never_revive_an_old_export(self) -> None:
        """Invalidate downloads before parsing changed bounds and ignore disabled invalid text"""
        # ARRANGE
        app = self._set_fixed_axis(self.ready())
        app.button(key='generate').click().run()
        # ACT
        app.text_input(key='axis_max_шт.').set_value('NaN').run()
        # ASSERT
        self.assert_no_export(app)
        self.assertTrue(app.button(key='generate').disabled)
        app.checkbox(key='axis_fixed_шт.').set_value(False).run()
        self.assertFalse(app.button(key='generate').disabled)
        self.assertFalse(app.error)
        app.button(key='generate').click().run()
        app.checkbox(key='axis_fixed_шт.').set_value(True).run()
        self.assert_no_export(app)
        app.text_input(key='axis_max_шт.').set_value('200').run()
        self.assertFalse(app.button(key='generate').disabled)
        self.assert_no_export(app)

    def test_all_outliers_block_generation_until_manual_correction(self) -> None:
        """Keep every plan and actual violation visible without expanding the scale silently"""
        # ARRANGE
        app = self._set_fixed_axis(self.ready())
        raw = 'показатель;план;факт;единица\nЗаявки;201;230;шт.\n'.encode()
        # ACT
        self.upload(app, raw)
        app.text_input(key='period').set_value('Неделя 2').run()
        # ASSERT
        warnings = '\n'.join(message.value for message in app.warning)
        self.assertIn('План: 201', warnings)
        self.assertIn('Факт: 230', warnings)
        self.assertIn('0…250', warnings)
        self.assertTrue(app.button(key='generate').disabled)
        self.assert_no_export(app)
        self.assertEqual(app.text_input(key='axis_max_шт.').value, '200')
        app.text_input(key='axis_max_шт.').set_value('250').run()
        self.assertFalse(app.warning)
        self.assertFalse(app.button(key='generate').disabled)
        app.button(key='generate').click().run()
        self.assertIsNotNone(app.session_state['pptx'])

    def test_hidden_invalid_settings_survive_replacement_without_blocking_other_units(self) -> None:
        """Restore raw settings after widget cleanup and preserve per-unit identity"""
        # ARRANGE
        app = self._set_fixed_axis(self.ready(), maximum='NaN')
        rubles = 'показатель;план;факт;единица\nВыручка;1;2;руб.\n'.encode()
        both = 'показатель;план;факт;единица\nВыручка;1;2;руб.\nЗаявки;120;135;ШТ.\n'.encode()
        # ACT
        self.upload(app, rubles)
        app.text_input(key='period').set_value('Неделя').run()
        # ASSERT
        self.assertFalse(app.button(key='generate').disabled)
        self.assertTrue(app.button(key='axis_save').disabled)
        self.assertEqual(app.session_state['axis_settings'].get('шт.'),
                         {'fixed': True, 'minimum': '0', 'maximum': 'NaN'})
        app.button(key='generate').click().run()
        self.upload(app, rubles.replace(b';1;', b';NaN;'))
        self.assert_no_export(app)
        self.upload(app, both)
        self.assertTrue(app.checkbox(key='axis_fixed_шт.').value)
        self.assertEqual(app.text_input(key='axis_max_шт.').value, 'NaN')
        self.assertFalse(app.checkbox(key='axis_fixed_руб.').value)
        self.assertTrue(app.button(key='generate').disabled)
        app.segmented_control(key='delimiter').set_value(',').run()
        app.segmented_control(key='delimiter').set_value(';').run()
        self.assertEqual(app.text_input(key='axis_max_шт.').value, 'NaN')

    def test_bookmark_restores_once_saves_explicitly_and_preserves_other_query_parameters(self) -> None:
        """Restore a fresh session without carrying CSV or overwriting unsaved edits on reruns"""
        # ARRANGE
        payload = '["report-slides-axis-v1",[["шт.","0","200"]]]'
        app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=15)
        app.query_params = {'scales': payload, 'keep': ['yes', 'again']}
        app.run()
        # ACT
        self.upload(app)
        app.text_input(key='period').set_value('Неделя').run()
        app.text_input(key='axis_max_шт.').set_value('250').run()
        app.run()
        # ASSERT
        self.assertEqual(app.text_input(key='axis_max_шт.').value, '250')
        self.assertEqual(app.query_params, {'scales': payload, 'keep': ['yes', 'again']})
        self.assertTrue(any(message.value == 'Изменения шкал не сохранены в адресе' for message in app.caption))
        app.button(key='axis_save').click().run()
        saved = '["report-slides-axis-v1",[["шт.","0","250"]]]'
        self.assertEqual(app.query_params, {'scales': saved, 'keep': ['yes', 'again']})
        another = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=15)
        another.query_params = dict(app.query_params)
        another.run()
        self.assertIsNone(another.session_state['metrics'])
        self.upload(another)
        self.assertTrue(another.checkbox(key='axis_fixed_шт.').value)
        self.assertEqual(another.text_input(key='axis_max_шт.').value, '250')
        self.assertEqual(another.text_input(key='period').value, '')
        another.button(key='axis_reset').click().run()
        self.assertEqual(another.query_params, {'keep': ['yes', 'again']})
        self.assertFalse(another.checkbox(key='axis_fixed_шт.').value)
        self.assertIsNotNone(another.session_state['metrics'])
        self.assertEqual(app.query_params, {'scales': saved, 'keep': ['yes', 'again']})

    def test_canonical_bookmark_comparison_differs_from_raw_export_invalidation(self) -> None:
        """Treat equivalent saved numbers equally while invalidating a changed PPTX input"""
        # ARRANGE
        app = self._set_fixed_axis(self.ready())
        app.button(key='axis_save').click().run()
        app.button(key='generate').click().run()
        # ACT
        app.text_input(key='axis_max_шт.').set_value('200,00').run()
        # ASSERT
        self.assert_no_export(app)
        self.assertFalse(any(message.value == 'Изменения шкал не сохранены в адресе' for message in app.caption))
        self.assertEqual(app.query_params, {'scales': '["report-slides-axis-v1",[["шт.","0","200"]]]'})
        app.checkbox(key='axis_fixed_шт.').set_value(False).run()
        app.button(key='axis_save').click().run()
        self.assertEqual(app.query_params, {})

    def test_malformed_bookmark_requires_explicit_reset(self) -> None:
        """Reject repeated or malformed URL input atomically and keep reset available before CSV"""
        # ARRANGE
        payloads = ('', '{}', ['bad', 'also bad'])
        # ACT
        for payload in payloads:
            app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=15)
            app.query_params = {'scales': payload, 'keep': 'yes'}
            app.run()
            # ASSERT
            self.assertFalse(app.exception)
            self.assertTrue(app.error)
            self.assertEqual(app.session_state['axis_settings'], {})
            self.upload(app)
            app.text_input(key='period').set_value('Неделя').run()
            self.assertTrue(app.button(key='generate').disabled)
            app.button(key='axis_reset').click().run()
            self.assertFalse(app.error)
            self.assertFalse(app.button(key='generate').disabled)
            self.assertEqual(app.query_params, {'keep': 'yes'})
            self.assertEqual(app.text_input(key='period').value, 'Неделя')

    def test_fixed_settings_work_without_network_and_escape_warning_text(self) -> None:
        """Preserve literal user text in warning Markdown and avoid external HTTP calls"""
        # ARRANGE
        raw = 'показатель;план;факт;единица\n``https://example.test/a``;201;230;**шт.**\n'.encode()
        app = self.upload(self.app(), raw)
        app.text_input(key='period').set_value('Неделя').run()
        # ACT
        with patch.object(socket.socket, 'connect', side_effect=AssertionError('External network')):
            self._set_fixed_axis(app, '**шт.**')
            app.button(key='axis_save').click().run()
            warnings = [message.value for message in app.warning]
            app.text_input(key='axis_max_**шт.**').set_value('250').run()
            app.button(key='generate').click().run()
            app.run()
        # ASSERT
        self.assertFalse(app.exception)
        self.assertIsNotNone(app.session_state['pptx'])
        self.assertTrue(any('``` Показатель «``https://example.test/a``», единица «**шт.**»' in message
                            for message in warnings))
        self.assertEqual(report_data.parse_axis_settings(app.query_params.get('scales')),
                         {'**шт.**': {'fixed': True, 'minimum': '0', 'maximum': '200'}})


if __name__ == '__main__':
    unittest.main()
