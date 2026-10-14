import argparse
from io import BytesIO
import logging
import os
from pathlib import Path
import tempfile
from urllib.parse import urlparse

from playwright.sync_api import sync_playwright, expect
from pptx import Presentation

expect.set_options(timeout=15000)


logging.basicConfig(level=logging.INFO)
logger = logging.getLogger('report_slides.browser_smoke')
parser = argparse.ArgumentParser(description='Проверка Report Slides в браузере')
parser.add_argument('--output', type=Path, help='Каталог снимков; по умолчанию создаётся временный каталог')
arguments = parser.parse_args()
output = (arguments.output if arguments.output is not None else
          Path(tempfile.mkdtemp(prefix='report-slides-screenshots-')))
output.mkdir(parents=True, exist_ok=True)
root = Path(__file__).resolve().parents[1]
sample = (root / 'samples/weekly-report.csv').read_bytes()
external = []
errors = []
browser_path = os.environ.get('REPORT_SLIDES_BROWSER')
with tempfile.TemporaryDirectory(prefix='report-slides-browser-') as temporary, sync_playwright() as driver:
    downloads = Path(temporary)
    browser = driver.chromium.launch(executable_path=browser_path, headless=True, args=['--no-sandbox'])
    context = browser.new_context(accept_downloads=True, viewport={'width': 1440, 'height': 1050})
    context.on('request', lambda request: external.append(request.url)
               if urlparse(request.url).hostname not in ('127.0.0.1', None) else None)
    page = context.new_page()
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.goto('http://127.0.0.1:8504')
    expect(page.get_by_role('heading', name='Report Slides', exact=True)).to_be_visible()
    generate = page.get_by_role('button', name='Создать презентацию', exact=True)
    expect(generate).to_be_disabled()
    expect(page.get_by_text('Период: заполните поле.', exact=True)).to_have_count(0)
    with page.expect_download() as received:
        page.get_by_role('button', name='Скачать пример CSV', exact=True).click()
    example = received.value
    assert example.suggested_filename == 'weekly-report.csv'
    example.save_as(downloads / 'weekly-report.csv')
    assert (downloads / 'weekly-report.csv').read_bytes() == sample
    upload = page.locator('input[type=file]')
    upload.set_input_files({'name': 'weekly.csv', 'mimeType': 'text/csv', 'buffer': sample})
    expect(page.get_by_text('5 слайдов: титул — 1, таблицы — 1, графики — 3.', exact=True)).to_be_visible()
    expect(page.get_by_text('Период: заполните поле.', exact=True)).to_have_count(0)
    page.get_by_label('Период', exact=True).fill('5–11 октября 2026')
    page.get_by_label('Период', exact=True).press('Tab')
    expect(generate).to_be_enabled()
    expect(page.locator('[data-testid=stApp]')).to_have_attribute('data-test-script-state', 'notRunning')
    page.get_by_label('Комментарий на титуле (необязательно)', exact=True).fill('Вымышленный пример для проверки.')
    page.get_by_label('Комментарий на титуле (необязательно)', exact=True).press('Tab')
    expect(generate).to_be_enabled()
    generate.click()
    pptx_button = page.get_by_role('button', name='Скачать PPTX', exact=True)
    expect(pptx_button).to_be_visible()
    with page.expect_download() as received:
        pptx_button.click()
    result = received.value
    assert result.suggested_filename == 'report-slides.pptx'
    result.save_as(downloads / 'first.pptx')
    first = (downloads / 'first.pptx').read_bytes()
    assert len(Presentation(BytesIO(first)).slides) == 5
    page.get_by_text(', — запятая', exact=True).click()
    expect(page.get_by_role('radio', name=', — запятая', exact=True)).to_be_checked()
    expect(pptx_button).to_have_count(0)
    expect(page.get_by_label('Период', exact=True)).to_have_value('5–11 октября 2026')
    expect(page.get_by_label('Комментарий на титуле (необязательно)', exact=True)).to_have_value(
        'Вымышленный пример для проверки.'
    )
    page.get_by_text('; — точка с запятой', exact=True).click()
    expect(page.get_by_role('radio', name='; — точка с запятой', exact=True)).to_be_checked()
    expect(generate).to_be_enabled()
    expect(pptx_button).to_have_count(0)
    generate.click()
    expect(pptx_button).to_be_visible()
    page.screenshot(path=str(output / 'browser-report.png'), full_page=True)
    page.get_by_role('heading', name='Report Slides', exact=True).scroll_into_view_if_needed()
    page.screenshot(path=str(output / 'browser-upload.png'), full_page=True)
    page.get_by_role('button', name='Загрузить новый CSV', exact=True).click()
    expect(pptx_button).to_have_count(0)
    expect(page.locator('[data-testid=stDataFrame]')).to_have_count(0)
    upload.set_input_files({'name': 'weekly.csv', 'mimeType': 'text/csv', 'buffer': sample.replace(b'120', b'121')})
    expect(page.get_by_label('Период', exact=True)).to_have_value('')
    expect(page.get_by_label('Комментарий на титуле (необязательно)', exact=True)).to_have_value('')
    expect(pptx_button).to_have_count(0)
    expect(generate).to_be_disabled()
    expect(page.locator('[data-testid=stApp]')).to_have_attribute('data-test-script-state', 'notRunning')
    expect(page.locator('[data-testid=stDataFrame]')).to_be_visible()
    expect(page.locator('[data-testid=stApp]')).to_have_attribute('data-test-script-state', 'notRunning')
    page.get_by_label('Период', exact=True).fill('12–18 октября 2026')
    page.get_by_label('Период', exact=True).press('Tab')
    expect(generate).to_be_enabled()
    generate.focus()
    page.keyboard.press('Enter')
    expect(pptx_button).to_be_visible()
    with page.expect_download() as received:
        pptx_button.click()
    received.value.save_as(downloads / 'updated.pptx')
    prs = Presentation(str(downloads / 'updated.pptx'))
    table = next(shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table)
    assert table.cell(1, 2).text == '121'
    page.get_by_role('button', name='Загрузить новый CSV', exact=True).click()
    expect(pptx_button).to_have_count(0)
    upload.set_input_files({'name': 'weekly.csv', 'mimeType': 'text/csv', 'buffer': sample.replace(b'120', b'NaN')})
    expect(page.get_by_text('Строка 2, поле «план»:', exact=False)).to_be_visible()
    expect(page.locator('[data-testid=stDataFrame]')).to_have_count(0)
    expect(pptx_button).to_have_count(0)
    page.get_by_text('Строка 2, поле «план»:', exact=False).scroll_into_view_if_needed()
    page.screenshot(path=str(output / 'browser-error.png'), full_page=True)
    page.reload()
    expect(page.get_by_role('heading', name='Report Slides', exact=True)).to_be_visible()
    expect(generate).to_be_disabled()
    expect(page.locator('[data-testid=stDataFrame]')).to_have_count(0)
    assert not external, external
    assert not errors, errors
    logger.info('Browser passed: downloads, CSV, PPTX, replacement, keyboard, refresh; screenshots: %s', output)
    browser.close()
